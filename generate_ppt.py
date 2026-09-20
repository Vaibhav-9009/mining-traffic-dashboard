from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

# Slide 1: Title
slide_1 = prs.slides.add_slide(prs.slide_layouts[0])
title_1 = slide_1.shapes.title
subtitle_1 = slide_1.placeholders[1]
title_1.text = "AI-Enabled Smart Barrier System"
subtitle_1.text = "Control Tower for Mixed-Fleet Traffic Management\nSystem Architecture & Status Update"

# Slide 2: The Core Problem & Solution
slide_2 = prs.slides.add_slide(prs.slide_layouts[1])
slide_2.shapes.title.text = "The Challenge & Our Solution"
tf = slide_2.placeholders[1].text_frame
tf.text = "The Challenge:"
p = tf.add_paragraph()
p.text = "- Mixed-fleet interactions (Haul Trucks vs Light Vehicles) are the leading cause of mining fatalities."
p.level = 1
p = tf.add_paragraph()
p.text = "- Traditional barriers are manual and slow."
p.level = 1
tf.add_paragraph().text = "Our Solution:"
p = tf.add_paragraph()
p.text = "- AI-powered control tower integrating Edge Computer Vision and GPS Telematics."
p.level = 1
p = tf.add_paragraph()
p.text = "- Automated rules engine that physically blocks unsafe vehicle interactions."
p.level = 1

# Slide 3: Live Dashboard Features
slide_3 = prs.slides.add_slide(prs.slide_layouts[1])
slide_3.shapes.title.text = "Live Dashboard Capabilities"
tf = slide_3.placeholders[1].text_frame
tf.text = "Event-Driven React Architecture:"
p = tf.add_paragraph()
p.text = "- Sub-second WebSockets delivery (No page refreshes)."
p.level = 1
p = tf.add_paragraph()
p.text = "- Real-Time Traffic Analytics: Recharts plotting traffic density and fleet composition."
p.level = 1
p = tf.add_paragraph()
p.text = "- Alert Center & Event History: Complete chronological auditing."
p.level = 1
p = tf.add_paragraph()
p.text = "- Human-in-the-Loop Override: 1-click barrier manual command."
p.level = 1

# Slide 4: Telematics & Geofencing Engine
slide_4 = prs.slides.add_slide(prs.slide_layouts[1])
slide_4.shapes.title.text = "Telematics (WheelsEye) & Geofencing"
tf = slide_4.placeholders[1].text_frame
tf.text = "Bypassing Camera Limitations:"
p = tf.add_paragraph()
p.text = "- We built a POST /api/telematics/webhook to ingest 3rd party GPS APIs."
p.level = 1
p = tf.add_paragraph()
p.text = "- Custom Geofencing Math: Translates standard Lat/Lng into designated Mining Zones (Zone A, Intersection, etc)."
p.level = 1
p = tf.add_paragraph()
p.text = "- Fuses broad fleet GPS data with highly localized AI gate tracking."
p.level = 1

# Slide 5: High-Density Simulation Results
slide_5 = prs.slides.add_slide(prs.slide_layouts[1])
slide_5.shapes.title.text = "Heavy Load & Rules Engine Validation"
tf = slide_5.placeholders[1].text_frame
tf.text = "Testing the Deterministic Engine:"
p = tf.add_paragraph()
p.text = "- Generated a massive 25-vehicle concurrent shift-change dataset."
p.level = 1
p = tf.add_paragraph()
p.text = "- Results: System dynamically flipped overcrowded zones into CONFLICT state."
p.level = 1
p = tf.add_paragraph()
p.text = "- Results: Overspeed alerts successfully mapped to correct offending tracking IDs."
p.level = 1
p = tf.add_paragraph()
p.text = "- Dashboard UI maintained high frame-rate rendering under heavy data load."
p.level = 1

# Save Presentation
prs.save("Project_Presentation_V2.pptx")
print("PPTX generated successfully as V2.")
